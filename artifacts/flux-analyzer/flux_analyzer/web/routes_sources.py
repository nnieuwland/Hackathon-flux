"""File upload, explicit grid selection and synthetic example sources."""
from pathlib import Path
import uuid
import zipfile
from flask import Blueprint, jsonify, request
from werkzeug.utils import secure_filename
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
from engine.ingest import list_sheets, preview, detect_header_row, load_source
from engine.money import clean
from config import WORKSPACE
from .state import get_project, save_project, public_state, source_path, load_defined, bind_sources

bp = Blueprint("sources", __name__)

def find_source(state: dict, sid: str) -> dict:
    """Find only sources belonging to the current session project."""
    source = next((s for s in state["sources"] if s["id"] == sid), None)
    if not source:
        raise ValueError("Source not found in this project.")
    return source

def add_file(state: dict, path: Path, name: str) -> dict:
    """Register a validated CSV or workbook under an opaque ID."""
    sheets = list_sheets(path)
    entry = {"id": uuid.uuid4().hex, "name": name, "stored": path.name, "size": path.stat().st_size, "sheets": sheets, "defined": False, "role": "primary", "columns": [], "n_rows": 0, "definition": {"sheet": sheets[0], "header_row": 1, "start_col": "A", "end_col": None, "role": "primary", "name": Path(name).stem, "combine_sheets": [], "tag_column": None}}
    state["sources"].append(entry)
    state["controls"] = []
    return entry

@bp.post("/sources/upload")
def upload() -> object:
    """Upload CSV/XLSX; reject executable, macro, malformed or oversized content."""
    state = get_project()
    files = request.files.getlist("files")
    if not files:
        raise ValueError("Choose at least one CSV or XLSX file.")
    directory = WORKSPACE / "uploads" / state["project_id"]
    directory.mkdir(parents=True, exist_ok=True)
    added_paths = []
    try:
        for file in files:
            name = secure_filename(file.filename or "")
            extension = Path(name).suffix.lower()
            if extension not in (".csv", ".xlsx"):
                raise ValueError("Only CSV and XLSX files are supported.")
            path = directory / f"{uuid.uuid4().hex}{extension}"
            file.save(path)
            added_paths.append(path)
            if extension == ".xlsx":
                with zipfile.ZipFile(path) as archive:
                    if sum(i.file_size for i in archive.infolist()) > 1024 * 1024 * 1024:
                        raise ValueError("This workbook expands beyond the 1 GB safety limit.")
            add_file(state, path, name)
    except Exception:
        for path in added_paths:
            path.unlink(missing_ok=True)
        raise
    save_project(state)
    return jsonify(public_state(state))

@bp.get("/sources/<sid>/preview")
def source_preview(sid: str) -> object:
    """Raw 25x30 preview with a separate parsed preview of user selections."""
    state = get_project()
    source = find_source(state, sid)
    path = source_path(state["project_id"], source)
    sheet = request.args.get("sheet") or source["sheets"][0]
    if sheet not in source["sheets"]:
        raise ValueError("Choose a sheet from this workbook.")
    grid = preview(path, sheet)
    header = int(request.args.get("header_row", detect_header_row(grid)))
    start = request.args.get("start_col", "A")
    end = request.args.get("end_col") or None
    parsed_error = None
    try:
        frame = load_source(path, sheet, header, start, end)
        parsed = {"columns": frame.columns.tolist(), "rows": clean(frame.head(5).fillna("").values.tolist()), "n_rows": len(frame)}
    except (ValueError, IndexError) as exc:
        parsed, parsed_error = {"columns": [], "rows": [], "n_rows": 0}, str(exc)
    width = max((len(row) for row in grid), default=0)
    return jsonify({"grid": grid, "suggested_header": detect_header_row(grid), "letters": [get_column_letter(i + 1) for i in range(width)], "sheets": source["sheets"], "parsed": parsed, "error": parsed_error})

@bp.post("/sources/<sid>/define")
def define_source(sid: str) -> object:
    """Persist only an explicit, valid file/sheet/row/column selection."""
    state, data = get_project(), request.get_json()
    source = find_source(state, sid)
    role = data.get("role", "primary")
    if role not in ("primary", "lookup", "forecast"):
        raise ValueError("Choose primary, lookup, or forecast as the source role.")
    sheet = data.get("sheet") or source["sheets"][0]
    combined = data.get("combine_sheets") or []
    if combined:
        combined = list(dict.fromkeys([sheet, *combined]))
    if sheet not in source["sheets"] or any(s not in source["sheets"] for s in combined):
        raise ValueError("Select sheets that exist in this workbook.")
    name = str(data.get("name") or Path(source["name"]).stem).strip()[:100]
    source["definition"] = {"sheet": sheet, "header_row": int(data.get("header_row", 1)), "start_col": str(data.get("start_col", "A")).upper(), "end_col": data.get("end_col") or None, "role": role, "name": name, "combine_sheets": combined, "tag_column": data.get("tag_column") or None}
    frame = load_defined(state["project_id"], source)
    source.update(defined=True, role=role, columns=list(frame.columns), n_rows=len(frame))
    state["controls"] = []
    bind_sources(state)
    from .state import auto_map_sources
    auto_map_sources(state)
    save_project(state)
    return jsonify(public_state(state))

@bp.delete("/sources/<sid>")
def remove_source(sid: str) -> object:
    """Remove a source from this project; retain run audit artifacts."""
    state = get_project()
    source = find_source(state, sid)
    state["sources"].remove(source)
    source_path(state["project_id"], source).unlink(missing_ok=True)
    state["sample_data"] = any(item.get("sample_data") for item in state["sources"])
    bind_sources(state)
    state["controls"] = []
    save_project(state)
    return jsonify(public_state(state))

@bp.post("/sources/demo")
def demo_sources() -> object:
    """Generate explicit synthetic ledgers, not fake UI metrics or results."""
    state = get_project()
    if any(s["name"] == "Sample_GL_2026.csv" for s in state["sources"]):
        state["sample_data"] = True
        save_project(state)
        return jsonify(public_state(state))
    import csv
    directory = WORKSPACE / "uploads" / state["project_id"]
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{uuid.uuid4().hex}.csv"
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Period_Name", "Functional_Net", "Country", "Account", "Region", "Po_Number", "FAH_Line_Description", "FAH_Event_Type", "Accounted_Dr", "Accounted_Cr"])
        countries = [("United Kingdom", 950000, 1280000, "Data centre cooling maintenance"), ("Netherlands", 740000, 540000, "Preventive maintenance contract"), ("Germany", 620000, 910000, "Emergency generator servicing"), ("France", 490000, 520000, "Electrical inspection"), ("Spain", 360000, 410000, "Fire safety maintenance")]
        for quarter, months in enumerate((["JAN-26", "FEB-26", "MAR-26"], ["APR-26", "MAY-26", "JUN-26"])):
            for ci, (country, b, c, text) in enumerate(countries):
                for mi, month in enumerate(months):
                    amount = (b if quarter == 0 else c) + (mi - 1) * 2100
                    writer.writerow([month, amount, country, "54100", "EMEA", f"70000{ci + 1}", text, "Operational", amount, 0])
        for month, amount, event in [("JAN-26", "1835.16", "Accrual - Current"), ("JAN-26", "-1813.59", "Accrual - Reversal"), ("APR-26", "1804.68", "Accrual - Current"), ("APR-26", "-1847.66", "Accrual - Reversal")]:
            writer.writerow([month, amount, "Netherlands", "54100", "EMEA", "800001", "Cooling accrual — synthetic golden rule example", event, amount if float(amount) > 0 else 0, -float(amount) if float(amount) < 0 else 0])
    sample = add_file(state, path, "Sample_GL_2026.csv")
    sample["sample_data"] = True
    book = Workbook()
    ws = book.active
    ws.title = "Open PO Report "
    ws.append(["PO Num", "GAO Comment"])
    ws.append(["800001", "GAO booked"])
    ws.append(["700001", "Open"])
    lookup_path = directory / f"{uuid.uuid4().hex}.xlsx"
    book.save(lookup_path)
    lookup = add_file(state, lookup_path, "Sample_Open_PO.xlsx")
    lookup["sample_data"] = True
    lookup["role"] = "lookup"
    lookup["definition"].update(role="lookup", name="open_po")
    state["sample_data"] = True
    save_project(state)
    return jsonify(public_state(state))
