"""Offline tests for the conversational journey: plan, steps, style, explanations, setups."""
import json
import time
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import create_app


@pytest.fixture
def api(tmp_path, monkeypatch):
    import web.routes_setups as setups
    # Saved setups are shared across users; keep test setups out of the real library.
    monkeypatch.setattr(setups, "WORKSPACE", tmp_path)
    app = create_app({"TESTING": True, "SECRET_KEY": "journey-test-only"})
    client = app.test_client()
    client.get("/")
    token = client.get("/state").get_json()["csrf_token"]
    headers = {"X-CSRF-Token": token}

    class Api:
        def get(self, path):
            response = client.get(path)
            assert response.status_code == 200, response.get_data(as_text=True)
            return response.get_json()

        def post(self, path, data=None, status=(200, 202)):
            response = client.post(path, json=data or {}, headers=headers)
            assert response.status_code in status, response.get_data(as_text=True)
            return response.get_json()

        def delete(self, path, data=None, status=(200,)):
            response = client.delete(path, json=data or {}, headers=headers)
            assert response.status_code in status, response.get_data(as_text=True)
            return response.get_json()

        def wait(self, job_id, expect_error=False):
            for _ in range(400):
                job = client.get(f"/jobs/{job_id}").get_json()
                if job["status"] == "done":
                    assert not expect_error, job
                    return job["result"]
                if job["status"] == "error":
                    assert expect_error, job.get("error")
                    return job
                time.sleep(.02)
            raise AssertionError("job timed out")

        raw = client
    return Api()


def demo_planned(api, description="Compare Q2 vs Q1, group by country, show the top 5 movements, threshold 100k"):
    api.post("/sources/demo")
    state = api.post("/sources/auto-define")
    assert all(source["defined"] for source in state["sources"])
    result = api.wait(api.post("/journey/plan", {"description": description})["job_id"])
    return result


def test_plan_from_description_applies_and_routes(api):
    result = demo_planned(api)
    assert result["applied"] and result["changes"]
    state = api.get("/state")
    spec = state["spec"]
    assert spec["comparison"]["current"]["months"] == ["2026-04", "2026-05", "2026-06"]
    assert spec["comparison"]["base"]["months"] == ["2026-01", "2026-02", "2026-03"]
    assert spec["drivers"]["top_n"] == 5
    assert spec["drivers"]["materiality"] == {"mode": "abs", "abs": 100000.0}
    view = next(v for v in spec["views"] if v["name"] == spec["drivers"]["view"])
    assert view["keys"] == ["Country"]
    journey = state["journey"]
    assert [step["id"] for step in journey["steps"]] == ["data", "compare", "include", "group", "important", "explain"]
    assert journey["ready_for_review"] and journey["first_open"] is None
    assert journey["steps"][0]["provenance"] in ("Set by AI", "Set automatically")


def test_plan_with_unknown_clause_warns_and_undo_restores(api):
    api.post("/sources/demo")
    api.post("/sources/auto-define")
    before = api.get("/state")["spec"]
    result = api.wait(api.post("/journey/plan", {"description": "Compare Q2 vs Q1; please do something impossible"})["job_id"])
    assert result["applied"]
    assert any("impossible" in warning for warning in result["warnings"])
    after = api.get("/state")
    assert after["spec"] != before
    restored = api.post("/spec/undo")
    assert restored["spec"]["comparison"] == before["comparison"]


def test_stale_ai_response_is_not_applied(api, monkeypatch):
    import web.routes_config as routes
    api.post("/sources/demo")
    api.post("/sources/auto-define")
    original = routes.invoke_agent
    def slow(envelope, **kwargs):
        response = original(envelope, **kwargs)
        # Someone changes the setup while the assistant is still working.
        api.post("/journey/simple/important", {"top_n": 7})
        return response
    monkeypatch.setattr(routes, "invoke_agent", slow)
    result = api.wait(api.post("/instruction", {"instruction": "Show the top 12 movements"})["job_id"])
    assert result["obsolete"] and not result["applied"]
    assert api.get("/state")["spec"]["drivers"]["top_n"] == 7


def test_instruction_outside_grammar_changes_nothing(api):
    before = api.get("/state")["spec"]
    result = api.wait(api.post("/instruction", {"instruction": "Rewrite the general ledger"})["job_id"])
    assert result["status"] == "cannot_do" and not result["applied"]
    assert api.get("/state")["spec"] == before


def test_protected_paths_are_rejected():
    from agent.contracts import check_patch, allowed_paths_for
    with pytest.raises(ValueError, match="protected"):
        check_patch([{"op": "replace", "path": "/sources", "value": {}}], allowed_paths_for(None))
    with pytest.raises(ValueError, match="outside"):
        check_patch([{"op": "replace", "path": "/drivers/top_n", "value": 3}], allowed_paths_for("explain"))
    with pytest.raises(ValueError, match="mandatory"):
        check_patch([], allowed_paths_for(None), {"controls": [{"id": "drivers_reconcile"}]}, {"controls": []})


def test_simple_step_forms_preserve_advanced_settings(api):
    demo_planned(api)
    state = api.get("/state")
    columns = [c["name"] for c in state["profile"]["columns"]]
    assert "Country" in columns
    api.post("/journey/simple/include", {"filters": [{"col": "Region", "values": ["EMEA"]}]})
    state = api.post("/journey/simple/group", {"keys": ["Country", "Account"]})
    view = next(v for v in state["spec"]["views"] if v["name"] == state["spec"]["drivers"]["view"])
    assert view["keys"] == ["Country", "Account"]
    assert state["spec"]["scope"] == [{"col": "Region", "op": "eq", "val": "EMEA"}]
    state = api.post("/journey/simple/important", {"threshold": 50000, "top_n": 4, "others_label": "Everything else"})
    assert state["spec"]["drivers"]["materiality"] == {"mode": "abs", "abs": 50000.0}
    assert state["spec"]["drivers"]["others_label"] == "Everything else"
    state = api.post("/journey/simple/explain", {"audience": "executive", "length": "short", "tone": "risk_focused", "structure_type": "split", "split_by": "Country"})
    profile = state["spec"]["commentary"]["style_profile"]
    assert profile["audience"] == "executive" and profile["structure"] == {"type": "split", "split_by": "Country"}
    bad = api.raw.post("/journey/simple/compare", json={"base_months": ["2026-01"], "current_months": ["2031-01"]}, headers={"X-CSRF-Token": api.get("/state")["csrf_token"]})
    assert bad.status_code == 400 and "not in your data" in bad.get_json()["error"]


def test_style_extraction_never_keeps_the_example(api):
    example = "- Germany up $1.2M in Q2 2026 due to Siemens generator servicing.\n- France adverse, partly offset by lower Spain costs."
    result = api.wait(api.post("/style/extract", {"example_text": example})["job_id"])
    assert result["applied"]
    state = api.get("/state")
    style = state["spec"]["commentary"]["style_profile"]["example_style"]
    dumped = json.dumps(state)
    for leaked in ("Siemens", "Germany", "1.2", "Q2 2026"):
        assert leaked not in json.dumps(style)
    assert "Siemens" not in dumped


def test_example_leak_check_catches_copied_facts():
    from agent.contracts import example_leaks
    issues = example_leaks({"summary_of_style": "Short bullets about Siemens.", "sample_phrases": ["Up $1.2M in May."]}, "We paid Siemens $1.2M in May.")
    assert len(issues) >= 2


def run_report(api):
    demo_planned(api)
    controls = api.post("/review/controls")["controls"]
    assert not any(c["status"] == "stop" for c in controls)
    run_id = api.wait(api.post("/run")["job_id"])["run_id"]
    return run_id


def test_explanation_statuses_and_human_confirmation(api):
    run_id = run_report(api)
    result = api.get(f"/results/{run_id}?format=json")
    statuses = {item["driver_id"]: item["status"] for item in result["explanations"]}
    # Deterministic placeholders are never treated as explanations.
    assert all(status in ("needs_explanation", "below_threshold") for status in statuses.values())
    draft = api.wait(api.post(f"/results/{run_id}/narrative")["job_id"])
    assert draft["applied"]
    result = api.get(f"/results/{run_id}?format=json")
    counts = result["explanation_counts"]
    assert counts["ai_review"] >= 1 and counts["confirmed"] == 0
    target = next(item for item in result["explanations"] if item["status"] == "ai_review")
    marked = api.post(f"/results/{run_id}/explanations/{target['driver_id']}", {"action": "save", "text": f"{target['delta_fmt']} — {target['label']}. [TO VERIFY] Ask the site lead.", "confirmed_by": "Ana"})
    assert next(i for i in marked["explanations"] if i["driver_id"] == target["driver_id"])["status"] == "needs_explanation"
    rejected = api.raw.post(f"/results/{run_id}/explanations/{target['driver_id']}", json={"action": "confirm", "confirmed_by": "Ana"}, headers={"X-CSRF-Token": api.get("/state")["csrf_token"]})
    assert rejected.status_code == 400
    saved = api.post(f"/results/{run_id}/explanations/{target['driver_id']}", {"action": "save", "text": f"{target['delta_fmt']} — {target['label']}: new maintenance contract signed.", "confirmed_by": "Ana"})
    item = next(i for i in saved["explanations"] if i["driver_id"] == target["driver_id"])
    assert item["status"] == "confirmed" and item["confirmed_by"] == "Ana"
    confirmed = api.post(f"/results/{run_id}/explanations/confirm-all", {"confirmed_by": "Ana"})
    assert confirmed["explanation_counts"]["ai_review"] == 0
    # A new AI draft must not overwrite confirmed text.
    api.wait(api.post(f"/results/{run_id}/narrative")["job_id"])
    after = api.get(f"/results/{run_id}?format=json")
    item = next(i for i in after["explanations"] if i["driver_id"] == target["driver_id"])
    assert item["status"] == "confirmed" and "new maintenance contract" in item["text"]


def test_excel_pack_has_formula_backed_primary_sheets(api):
    run_id = run_report(api)
    data = api.raw.get(f"/download/{run_id}/excel").data
    book = load_workbook(BytesIO(data))
    assert book.sheetnames[:7] == ["Summary", "Needs Explanation", "Commentary", "Drill-down", "Data checks", "Data", "Settings"]
    assert {"Movement detail", "Data controls", "FluxSpec", "Setup changes", "Run audit", "AI assistant log"} <= set(book.sheetnames)
    commentary = book["Commentary"]
    formulas = [cell.value for row in commentary.iter_rows(min_row=2) for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert any("SUMPRODUCT" in formula for formula in formulas)
    assert "BaseMonths" in book.defined_names and "CurrentMonths" in book.defined_names and "Threshold" in book.defined_names
    cached = load_workbook(BytesIO(data), data_only=True)
    checks = cached["Data checks"]
    statuses = [row[-1].value for row in checks.iter_rows(min_row=2) if row[-1].value in ("OK", "Check")]
    assert statuses and set(statuses) == {"OK"}


def test_saved_setups_are_versioned_and_conflicts_are_refused(api):
    demo_planned(api)
    created = api.post("/setups", {"name": "Country review", "description": "Quarterly by country", "saved_by": "Ana"})["setup"]
    assert created["current_version"] == 1
    duplicate = api.raw.post("/setups", json={"name": "country review"}, headers={"X-CSRF-Token": api.get("/state")["csrf_token"]})
    assert duplicate.status_code == 400
    api.post("/journey/simple/important", {"top_n": 9})
    second = api.post(f"/setups/{created['id']}/versions", {"expected_version": 1, "note": "Nine movements", "saved_by": "Ben"})["setup"]
    assert second["current_version"] == 2 and second["updated_by"] == "Ben"
    stale = api.raw.post(f"/setups/{created['id']}/versions", json={"expected_version": 1}, headers={"X-CSRF-Token": api.get("/state")["csrf_token"]})
    assert stale.status_code == 400 and "version 2" in stale.get_json()["error"]
    loaded = api.post(f"/setups/{created['id']}/load", {"version": 1})
    assert loaded["spec"]["drivers"]["top_n"] == 5 and loaded["loaded_setup"]["version"] == 1
    exported = api.raw.get(f"/setups/{created['id']}/export?version=2")
    body = json.loads(exported.data)
    assert body["spec"]["sources"] == {} and body["version"] == 2
    copy = api.post(f"/setups/{created['id']}/duplicate", {"name": "Copy"})["setup"]
    api.post(f"/setups/{copy['id']}/rename", {"name": "Renamed copy"})
    refused = api.raw.delete(f"/setups/{copy['id']}", json={"confirm_name": "wrong"}, headers={"X-CSRF-Token": api.get("/state")["csrf_token"]})
    assert refused.status_code == 400
    remaining = api.delete(f"/setups/{copy['id']}", {"confirm_name": "Renamed copy"})["setups"]
    assert all(item["id"] != copy["id"] for item in remaining)
    assert api.raw.get("/setups/..%2F..%2Fconfig").status_code in (400, 404)


def test_flow_pages_and_old_links_render(api):
    for path in ("/", "/flow/upload", "/flow/data", "/flow/results", "/step/results", "/step/sources", "/review"):
        response = api.raw.get(path)
        assert response.status_code == 200, path
        assert b"journey" in response.data
    assert api.raw.get("/flow/nope").status_code == 404


def _evaluate_sumproduct(book, sheet, formula):
    """Tiny evaluator for the pack's exact-match SUMPRODUCT shape, independent of Excel."""
    import re
    from openpyxl.utils import column_index_from_string
    data = book["Data"]
    def months(name):
        target = book.defined_names[name].attr_text.replace("$", "").split("!")[1]
        start, end = target.split(":")
        return {cell.value for row in book["Settings"][start:end] for cell in row}
    keys = [(column_index_from_string(col), sheet[f"{ref_col}{ref_row}"].value) for col, ref_col, ref_row in re.findall(r"\(Data!\$([A-Z]+)\$2:\$[A-Z]+\$\d+=\$?([A-Z]+)\$?(\d+)\)", formula)]
    month_col, side = re.search(r"MATCH\(Data!\$([A-Z]+)\$2:\$[A-Z]+\$\d+,(Base|Current)Months", formula).groups()
    amount_col = re.search(r"\*Data!\$([A-Z]+)\$2:\$[A-Z]+\$\d+\)$", formula).group(1)
    wanted = months(f"{side}Months")
    total = 0.0
    for row in data.iter_rows(min_row=2, values_only=True):
        same = lambda a, b: str(a).casefold() == str(b).casefold()  # Excel's "=" ignores case
        if row[column_index_from_string(month_col) - 1] in wanted and all(same(row[col - 1], value) for col, value in keys):
            total += row[column_index_from_string(amount_col) - 1]
    return total


def test_excel_formulas_reproduce_engine_values(api):
    run_id = run_report(api)
    data = api.raw.get(f"/download/{run_id}/excel").data
    book, cached = load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)
    result = api.get(f"/results/{run_id}?format=json")
    sheet, values = book["Commentary"], cached["Commentary"]
    headers = [cell.value for cell in sheet[1]]
    checked = 0
    for row in range(2, sheet.max_row + 1):
        for name in ("Base", "Current"):
            col = headers.index(name) + 1
            formula = sheet.cell(row, col).value
            if isinstance(formula, str) and "SUMPRODUCT" in formula:
                assert abs(_evaluate_sumproduct(book, sheet, formula) - values.cell(row, col).value) < 0.01
                checked += 1
    assert checked >= 2
    totals = cached["Summary"]
    assert abs(totals["A5"].value - result["base"]) < 0.01 and abs(totals["B5"].value - result["current"]) < 0.01
    assert abs(_evaluate_sumproduct(book, book["Summary"], book["Summary"]["A5"].value) - result["base"]) < 0.01
    assert abs(_evaluate_sumproduct(book, book["Summary"], book["Summary"]["B5"].value) - result["current"]) < 0.01


def test_excel_falls_back_to_engine_values_when_formulas_cannot_be_exact(api):
    demo_planned(api)
    spec = api.get("/state")["spec"]
    spec["comparison"]["base"]["filters"] = [{"col": "Country", "op": "ne", "val": "Spain"}]
    api.post("/config/update", {"spec": spec})
    run_id = api.wait(api.post("/run")["job_id"])["run_id"]
    book = load_workbook(BytesIO(api.raw.get(f"/download/{run_id}/excel").data))
    text = " ".join(str(cell.value) for row in book["Settings"].iter_rows() for cell in row if cell.value)
    assert "Engine-calculated" in text and "extra filters" in text
    assert not any(isinstance(cell.value, str) and "SUMPRODUCT" in cell.value for row in book["Commentary"].iter_rows() for cell in row)


def test_stopping_the_assistant_discards_its_late_answer(api, monkeypatch):
    import web.routes_config as routes
    api.post("/sources/demo")
    api.post("/sources/auto-define")
    before = api.get("/state")["spec"]["drivers"]["top_n"]
    original = routes.invoke_agent
    def slow(envelope, **kwargs):
        response = original(envelope, **kwargs)
        # The user pressed "Continue without the assistant" while this was running.
        api.post("/journey/stop-assistant")
        return response
    monkeypatch.setattr(routes, "invoke_agent", slow)
    result = api.wait(api.post("/instruction", {"instruction": "Show the top 12 movements"})["job_id"])
    assert result["obsolete"] and not result["applied"]
    assert api.get("/state")["spec"]["drivers"]["top_n"] == before


def test_include_step_lists_column_values_and_rejects_unknown_columns(api):
    api.post("/sources/demo")
    api.post("/sources/auto-define")
    values = api.get("/journey/values?col=Country")
    assert values["n_distinct"] == len(values["values"]) > 1
    assert all(item["rows"] > 0 for item in values["values"])
    response = api.raw.get("/journey/values?col=__source")
    assert response.status_code == 400


def test_project_state_does_not_carry_the_export_rows(api):
    demo_planned(api)
    api.post("/journey/simple/compare", {"base_months": ["2026-01", "2026-02", "2026-03"], "current_months": ["2026-04", "2026-05", "2026-06"]})
    api.post("/journey/simple/group", {"keys": ["Country"]})
    run_id = api.wait(api.post("/run")["job_id"])["run_id"]
    state = api.get("/state")
    assert state["result"]["id"] == run_id and "export_data" not in state["result"]
    assert "export_data" not in api.get(f"/results/{run_id}?format=json")


def test_simple_mode_assets_and_plain_month_wording(api):
    page = api.raw.get("/flow/upload").get_data(as_text=True)
    assert "/static/journey.js" in page and "/static/journey.css" in page and "__JOURNEY__" in page
    assert page.index("/static/app.js") < page.index("/static/journey.js") < page.index("alpine.min.js")
    from web.journey import _month_range
    assert _month_range(["2026-04", "2026-05", "2026-06"]) == "Apr 2026 to Jun 2026"
    assert _month_range(["2026-01", "2026-03"]) == "Jan 2026 and Mar 2026"


def _grouping(state):
    spec = state["spec"]
    return next((view["keys"] for view in spec.get("views", []) if view["name"] == spec.get("drivers", {}).get("view")), [])


@pytest.mark.parametrize("description, base, current, keys", [
    ("Compare the latest quarter with the previous quarter by country", ["2026-01", "2026-02", "2026-03"], ["2026-04", "2026-05", "2026-06"], ["Country"]),
    ("June vs May by account, threshold 50k, top 15", ["2026-05"], ["2026-06"], ["Account"]),
    ("Compare Q2 with Q1 by region. Explain movements over 100k, top 10.", ["2026-01", "2026-02", "2026-03"], ["2026-04", "2026-05", "2026-06"], ["Region"]),
])
def test_plan_uses_every_part_of_the_description(api, description, base, current, keys):
    result = demo_planned(api, description)
    assert result["applied"] and not any("could not use" in warning for warning in result.get("warnings", []))
    state = api.get("/state")
    assert state["spec"]["comparison"]["base"]["months"] == base
    assert state["spec"]["comparison"]["current"]["months"] == current
    assert _grouping(state) == keys
    assert all(step["status"] == "ready" for step in state["journey"]["steps"])


def test_plan_flags_the_step_whose_wording_was_not_understood(api):
    demo_planned(api, "Compare Q2 with Q1 by colour")
    state = api.get("/state")
    group = next(step for step in state["journey"]["steps"] if step["id"] == "group")
    assert group["status"] == "needs_input" and _grouping(state) == []
    result = api.wait(api.post("/journey/plan", {"description": "Compare Q2 with Q1 by country and exclude intercompany"})["job_id"])
    assert any("exclude intercompany" in warning for warning in result["warnings"])
    include = next(step for step in api.get("/state")["journey"]["steps"] if step["id"] == "include")
    assert include["status"] == "needs_input"


def test_partially_understood_clauses_are_refused_not_half_applied():
    from agent.mock_instructions import clause_patch
    months = ["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"]
    assert clause_patch({}, "Compare the latest quarter with the same quarter last year", [], months) is None
    assert clause_patch({}, "threshold 50k for EMEA only", [], months) is None
    assert clause_patch({}, "short commentary about Germany", [], months) is None
    assert clause_patch({}, "Compare the latest quarter with the previous quarter", [], months) is not None


def test_style_extraction_rejects_copied_names_and_claims():
    from agent.contracts import example_leaks
    example = "Siemens drove the increase after the contract renegotiation was delayed.\n- Cooling accrual reversed."
    assert example_leaks({"summary_of_style": "Short bullets.", "sample_phrases": ["Siemens drove the increase."]}, example)
    assert example_leaks({"summary_of_style": "Notes that the contract renegotiation was delayed.", "sample_phrases": []}, example)
    safe = {"summary_of_style": "Bulleted, one movement per line, short sentences, labels movements favourable or adverse, separates timing effects.",
            "sample_phrases": ["The movement relates to [supported evidence].", "[Amount] — [movement]: [explanation].", "Partly offset by [related movement]."]}
    assert example_leaks(safe, example) == []


def test_excel_overview_groups_formats_and_checked_commentary(api):
    from pathlib import Path
    from engine.excel_pack import GROUP_COLOURS, PRIMARY
    run_id = run_report(api)
    api.wait(api.post(f"/results/{run_id}/narrative")["job_id"])
    result = api.get(f"/results/{run_id}?format=json")
    target = next(item for item in result["explanations"] if item["status"] == "ai_review")
    api.post(f"/results/{run_id}/explanations/{target['driver_id']}", {"action": "save", "text": f"{target['delta_fmt']} — {target['label']}: new maintenance contract signed.", "confirmed_by": "Ana"})
    book = load_workbook(BytesIO(api.raw.get(f"/download/{run_id}/excel").data))

    # Order: review and evidence sheets first, then supporting detail, then the audit trail, ending with the setup.
    names = book.sheetnames
    assert names[:len(PRIMARY)] == PRIMARY and names[-1] == "FluxSpec"
    colour = {name: (book[name].sheet_properties.tabColor.rgb[-6:] if book[name].sheet_properties.tabColor else None) for name in names}
    groups = {value.lstrip("#").upper(): key for key, value in GROUP_COLOURS.items()}
    order = [groups[colour[name]] for name in names]
    assert order == sorted(order, key=["Review", "Evidence and reconciliation", "Supporting detail", "Audit trail"].index)
    assert names.index("Source lines") < names.index("Data controls") < names.index("Setup changes")

    # Overview content: comparison, totals, status, key drivers with checked commentary, linked contents.
    summary = book["Summary"]
    cells = {cell.coordinate: cell for row in summary.iter_rows() for cell in row if cell.value is not None}
    text = " ".join(str(cell.value) for cell in cells.values())
    assert summary["A1"].value == result["name"] and "compared with" in summary["A2"].value and "Apr 2026" in summary["A2"].value
    for heading in ("Base total", "Movement %", "Explanation status", "Headline", "Key drivers", "What is in this workbook", "About these numbers"):
        assert heading in text, heading
    driver_row = next(cell.row for cell in cells.values() if cell.value == "Movement" and cell.column == 1)
    first = summary.cell(driver_row + 1, 1).value
    assert first == result["drivers"][0]["label"]
    assert summary.cell(driver_row + 1, 4).value.startswith("='Commentary'!")
    confirmed_row = next(row for row in range(driver_row + 1, driver_row + 12) if summary.cell(row, 1).value == target["label"])
    assert summary.cell(confirmed_row, 6).value == "Confirmed" and "maintenance contract" in summary.cell(confirmed_row, 7).value
    links = {cell.hyperlink.location for cell in cells.values() if cell.hyperlink}
    assert "'Needs Explanation'!A1" in links and "'FluxSpec'!A1" in links and "'Data checks'!A1" in links

    # Consistent financial formats.
    assert "#,##0.00" in summary["A5"].number_format and summary["D5"].number_format.startswith("0.0%")
    commentary = book["Commentary"]
    headers = [cell.value for cell in commentary[1]]
    assert commentary.cell(2, headers.index("Change") + 1).number_format.startswith("#,##0.00")
    assert commentary.cell(2, headers.index("Change %") + 1).number_format.startswith("0.0%")
    confirmed = next(row for row in range(2, commentary.max_row + 1) if commentary.cell(row, 1).value == target["label"])
    assert commentary.cell(confirmed, headers.index("Confirmed by") + 1).value == "Ana"
    assert commentary.cell(confirmed, headers.index("Confirmed (UTC)") + 1).is_date
    assert book["Data"].cell(2, book["Data"].max_column).number_format.startswith("#,##0.00")
    created = next(row for row in book["Settings"].iter_rows() if row[0].value == "Created (UTC)")[1]
    assert created.is_date
    variance = book[next(name for name in names if name.startswith("Variance"))]
    vheaders = [cell.value for cell in variance[1]]
    assert "Change %" in vheaders and "amount_delta_pct" not in vheaders
    pct = variance.cell(2, vheaders.index("Change %") + 1)
    assert pct.number_format.startswith("0.0%") and abs(pct.value) < 10
    detail = book["Movement detail"]
    dheaders = [cell.value for cell in detail[1]]
    assert detail.cell(2, dheaders.index("Lines") + 1).number_format == "#,##0"
    log = book["Explanation log"]
    assert log["A2"].is_date and any(cell.value == "Ana" for cell in log[2])

    # Word and Markdown are gone from both results pages.
    templates = Path(__file__).resolve().parent.parent / "web" / "templates"
    for page in ("index.html", "journey/simple.html"):
        html = (templates / page).read_text()
        assert "/docx" not in html and "/markdown" not in html and "/excel" in html and "/spec" in html


def _run_regions(api, regions):
    """Upload a small file whose grouping values are chosen by the test, run it and return the workbook."""
    rows = "".join(f"{month},{amount},{region}\n" for region, (base, current) in regions.items() for month, amount in (("APR-26", base), ("MAY-26", current)))
    data = ("Period,Amount,Region\n" + rows).encode()
    uploaded = api.raw.post("/sources/upload", data={"files": (BytesIO(data), "regions.csv")}, content_type="multipart/form-data", headers={"X-CSRF-Token": api.get("/state")["csrf_token"]})
    source = uploaded.get_json()["sources"][0]
    api.post("/sources/" + source["id"] + "/define", {"name": "regions", "role": "primary", "header_row": 1, "start_col": "A"})
    settings = {
        "roles": {"value": "Amount", "period": "Period", "scenario": None, "debit": None, "credit": None, "keys": [], "dimensions": ["Region"], "driver_text": []},
        "comparison": {"type": "MoM", "base": {"months": ["2026-04"]}, "current": {"months": ["2026-05"]}},
        "measures": {"amount": {"agg": "sum", "col": "{roles.value}"}},
        "views": [{"name": "regions", "keys": ["Region"], "measures": ["amount"], "compare": True}],
        "drivers": {"view": "regions", "measure": "amount", "materiality": {"mode": "abs", "abs": 0}, "top_n": 15, "others_label": "Other items"},
    }
    api.post("/spec/patch", {"patch": [{"op": "add", "path": "/" + key, "value": value} for key, value in settings.items()], "reason": "Region test"})
    run_id = api.wait(api.post("/run")["job_id"])["run_id"]
    data = api.raw.get(f"/download/{run_id}/excel").data
    return api.get(f"/results/{run_id}?format=json"), load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _recalculated(book, result):
    """Each movement's Base/Current as Excel would recalculate them from the Data sheet."""
    sheet = book["Commentary"]
    headers = [cell.value for cell in sheet[1]]
    amounts = {}
    for row in range(2, sheet.max_row + 1):
        formulas = [sheet.cell(row, headers.index(name) + 1).value for name in ("Base", "Current")]
        if all(isinstance(formula, str) and "SUMPRODUCT" in formula for formula in formulas):
            amounts[sheet.cell(row, 1).value] = tuple(_evaluate_sumproduct(book, sheet, formula) for formula in formulas)
    return amounts


def test_formula_like_region_names_recalculate_to_engine_values(api):
    result, book, _ = _run_regions(api, {"=North": (100, 180), "South": (50, 40), "@East": (30, 75), "+West": (10, 5)})
    amounts = _recalculated(book, result)
    assert len(amounts) == 4
    engine = {driver["label"]: (driver["base"], driver["current"]) for driver in result["drivers"]}
    for label, (base, current) in engine.items():
        shown = "'" + label if label[0] in "=@+" else label
        assert amounts[shown] == (base, current), label


def test_region_names_that_collide_once_protected_fall_back_to_engine_values(api):
    result, book, cached = _run_regions(api, {"=North": (100, 180), "'=North": (1000, 1500), "South": (50, 40)})
    assert {driver["label"] for driver in result["drivers"]} >= {"=North", "'=North"}
    # Formulas would add both "=North" groups together, so the pack must not use them.
    assert not _recalculated(book, result)
    assert not any(isinstance(cell.value, str) and "SUMPRODUCT" in cell.value for sheet in book.worksheets for row in sheet.iter_rows() for cell in row)
    text = " ".join(str(cell.value) for row in book["Settings"].iter_rows() for cell in row if cell.value)
    assert "Engine-calculated" in text and "protected against formulas" in text
    sheet, values = book["Commentary"], cached["Commentary"]
    headers = [cell.value for cell in sheet[1]]
    by_row = {(values.cell(row, headers.index("Base") + 1).value, values.cell(row, headers.index("Current") + 1).value) for row in range(2, sheet.max_row + 1)}
    assert {(100, 180), (1000, 1500)} <= by_row
