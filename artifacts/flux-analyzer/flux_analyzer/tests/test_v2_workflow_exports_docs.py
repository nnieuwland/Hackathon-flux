"""Whole-request integration: user preset, run, line edits, audit and reuse."""
from io import BytesIO
import json
import time
from openpyxl import load_workbook
from app import create_app
from engine.spec import validate_spec, default_spec
from engine.audit import spec_hash


def test_user_created_preset_run_exports_history_and_docs(tmp_path, monkeypatch):
    import engine.presets as presets
    monkeypatch.setattr(presets, "WORKSPACE", tmp_path)
    client = create_app({"TESTING": True, "SECRET_KEY": "v2-workflow-test"}).test_client()
    initial = client.get("/state").get_json()
    headers = {"X-CSRF-Token": initial["csrf_token"]}
    def post(path, data):
        reply = client.post(path, json=data, headers=headers)
        assert reply.status_code in (200, 202), reply.get_json()
        return reply.get_json()
    def wait(job):
        for _ in range(200):
            status = client.get("/jobs/" + job["job_id"]).get_json()
            if status["status"] == "done":
                return status["result"]
            assert status["status"] != "error", status.get("error")
            time.sleep(.02)
        raise AssertionError("Background job did not finish")
    data = b"Period,Amount,Region,Text\nAPR-26,100,EMEA,Repairs\nAPR-26,50,APAC,Office\nMAY-26,150,EMEA,Repairs\nMAY-26,40,APAC,Office\n"
    uploaded = client.post("/sources/upload", data={"files": (BytesIO(data), "costs.csv")}, content_type="multipart/form-data", headers=headers)
    source = uploaded.get_json()["sources"][0]
    post("/sources/" + source["id"] + "/define", {"name": "costs", "role": "primary", "header_row": 1, "start_col": "A"})
    settings = {
        "roles": {"value": "Amount", "period": "Period", "scenario": None, "debit": None, "credit": None, "keys": [], "dimensions": ["Region"], "driver_text": ["Text"]},
        "comparison": {"type": "MoM", "base": {"months": ["2026-04"]}, "current": {"months": ["2026-05"]}},
        "measures": {"amount": {"agg": "sum", "col": "{roles.value}"}},
        "views": [{"name": "regions", "keys": ["Region"], "measures": ["amount"], "compare": True}],
        "drivers": {"view": "regions", "measure": "amount", "materiality": {"mode": "abs", "abs": 0}, "top_n": 15, "others_label": "Other items"},
        "classification": {"target_col": "Category", "source_cols": ["Text"], "fallback": "Other", "rules": [{"keyword": "Repairs", "category": "Maintenance", "match": "contains", "reviewed": True}]},
    }
    post("/spec/patch", {"patch": [{"op": "add", "path": "/" + key, "value": value} for key, value in settings.items()], "reason": "Created a user Opex configuration"})
    saved = post("/presets/save", {"id": "user_opex", "name": "User Opex", "owner": "Finance", "description": "Monthly operating costs", "guidance": "Explain region movements from checked evidence."})
    assert saved["state"]["spec"]["preset"]["id"] == "user_opex"
    preview = post("/classification/preview", {})
    assert preview["stats"]["row_coverage"] == 50
    for section in ("scope", "derive", "joins", "value_rules"):
        assert "n_rows" in post("/config/preview/" + section, {})
    run_id = wait(post("/run", {}))["run_id"]
    run = client.get(f"/results/{run_id}?format=json").get_json()
    assert run["delta"] == 40
    assert len(run["run_stamp"]["spec_sha256"]) == 64
    assert len(run["run_stamp"]["input_files"][0]["sha256"]) == 64
    assert run["change_log"] and run["commentary"]["lines"]
    generated = wait(post(f"/results/{run_id}/narrative", {}))
    assert generated["validation"]["ok"], generated
    regenerated = wait(post(f"/results/{run_id}/narrative", {"driver_id": "d1"}))
    assert len(regenerated["proposal"]["lines"]) == len(run["drivers"])
    edited = post(f"/results/{run_id}/commentary", {"text": "", "lines": regenerated["proposal"]["lines"]})
    assert edited["validation"]["ok"] and not edited["commentary"]["verified"]
    for kind in ("docx", "markdown"):
        retired = client.get(f"/download/{run_id}/{kind}")
        assert retired.status_code == 400 and "retired" in retired.get_json()["error"]
    for kind in ("spec", "excel"):
        download = client.get(f"/download/{run_id}/{kind}")
        assert download.status_code == 200
        if kind == "spec":
            exported = json.loads(download.data)
            assert validate_spec(exported)["run_stamp"]["run_id"] == run_id
        else:
            book = load_workbook(BytesIO(download.data))
            assert {"Classifications", "Unmatched lookups", "Setup changes", "Explanation log", "Run audit", "AI assistant log", "FluxSpec", "Source lines", "Movement detail"} <= set(book.sheetnames)
            audit = " ".join(str(cell.value) for row in book["Run audit"].iter_rows() for cell in row)
            assert run_id in audit and "SHA-256" in audit
    suggestion = wait(post("/instruction", {"instruction": "Set materiality to 250k"}))
    assert suggestion["applied"] and suggestion["changes"]
    base_spec = client.get("/state").get_json()["spec"]
    assert base_spec["drivers"]["materiality"]["abs"] == 250000
    preview = post("/config/draft/preview", {"spec": base_spec})
    draft = preview["spec"]
    changed = post("/config/draft/patch", {"base_spec": base_spec, "draft_spec": draft,
        "patch": [{"op": "replace", "path": "/flux_name", "value": "Reviewed monthly costs"}]})
    assert changed["spec"]["flux_name"] == "Reviewed monthly costs"
    assert client.get("/state").get_json()["spec"] == base_spec
    accepted = post("/config/draft/accept", {"base_spec": base_spec, "spec": changed["spec"], "reason": "Reviewed suggestion"})
    assert accepted["spec"]["flux_name"] == "Reviewed monthly costs"
    assert accepted["spec"]["drivers"]["materiality"]["abs"] == 250000
    reused = post(f"/runs/{run_id}/reuse", {})
    assert reused["spec"]["drivers"]["materiality"]["abs"] == 0
    assert reused["spec"]["sources"] and reused["runs"][0]["id"] == run_id
    answer = wait(post("/questions", {"question": "What does materiality mean here?", "run_id": run_id}))
    assert answer["proposal"]["answer"]
    assert client.get("/state").get_json()["last_answer"]["question"]
    from web.routes_docs import DOCUMENTS
    for identifier in DOCUMENTS:
        assert client.get("/help/" + identifier).status_code == 200
    assert client.get("/docs?doc=reference").status_code == 200
    assert client.get("/help/../../config").status_code == 400


def test_hash_is_stable_and_ignores_previous_run_stamp():
    settings = default_spec()
    assert spec_hash(settings) == spec_hash(dict(reversed(list(settings.items()))))
    assert spec_hash(settings) == spec_hash(dict(settings, run_stamp={"run_id": "prior"}))


def test_multiple_invalid_shapes_and_patterns_report_all_errors():
    import pytest
    settings = default_spec()
    settings.update(commentary=[], comparison="bad", classification="bad", derive=[{"name": "Key", "type": "regex_extract", "source": "Text", "pattern": "(", "strip_suffix": "["}])
    with pytest.raises(ValueError) as error:
        validate_spec(settings)
    for field in ("commentary", "comparison", "classification", "pattern", "strip_suffix"):
        assert field in str(error.value)
