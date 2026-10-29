"""CSV/XLSX ingestion, raw selection previews and content-keyed caching."""
import hashlib
from pathlib import Path
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

def list_sheets(path: str | Path) -> list[str]:
    """Return exact workbook sheet names, including trailing whitespace."""
    if Path(path).suffix.lower() == ".csv":
        return ["(csv)"]
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()

def preview(path: str | Path, sheet: str | None = None, n_rows: int = 25, n_cols: int = 30) -> list[list[str]]:
    """Read a bounded raw grid with no header interpretation."""
    if Path(path).suffix.lower() == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = []
            for i, row in enumerate(csv.reader(handle, dialect)):
                if i >= n_rows:
                    break
                rows.append([str(v) for v in row[:n_cols]])
            return rows
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = book[sheet or book.sheetnames[0]]
        return [[str(v) if v is not None else "" for v in row] for row in ws.iter_rows(max_row=min(ws.max_row or n_rows, n_rows), max_col=min(ws.max_column or n_cols, n_cols), values_only=True)]
    finally:
        book.close()

def detect_header_row(grid: list[list[str]]) -> int:
    """Suggest a 1-based header row; the user always confirms it."""
    scores = []
    for i, row in enumerate(grid[:15]):
        cells = [c for c in row if c.strip()]
        strings = sum(not c.replace(".", "", 1).replace("-", "", 1).isdigit() for c in cells)
        score = len(cells) + strings * 0.4 + len(set(cells)) * 0.2
        scores.append((score, -i))
    return -max(scores)[1] + 1 if scores else 1

def load_source(path: str | Path, sheet: str | None, header_row: int, start_col: str, end_col: str | None = None, cache_dir: Path | None = None) -> pd.DataFrame:
    """Apply an explicit selection, clean headers and cache trusted local frames."""
    path = Path(path)
    if header_row < 1:
        raise ValueError("Header row must be at least 1.")
    first = column_index_from_string(start_col.upper()) - 1
    last = column_index_from_string(end_col.upper()) if end_col else None
    if last is not None and last <= first:
        raise ValueError("Last column must be at or after the first column.")
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    digest.update(repr((sheet, header_row, start_col, end_col)).encode())
    key = digest.hexdigest()
    if cache_dir:
        cache_dir.mkdir(parents=True, exist_ok=True)
        for ext in ("parquet", "pkl"):
            cached = cache_dir / f"{key}.{ext}"
            if cached.exists():
                return pd.read_parquet(cached) if ext == "parquet" else pd.read_pickle(cached)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig") as handle:
            sample = handle.read(8192)
        try:
            sep = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
        except csv.Error:
            sep = ","
        kwargs = dict(header=None, skiprows=header_row - 1, low_memory=False, encoding="utf-8-sig", sep=sep)
        if path.stat().st_size > 500 * 1024 * 1024:
            frame = pd.concat(pd.read_csv(path, chunksize=100000, **kwargs), ignore_index=True)
        else:
            frame = pd.read_csv(path, **kwargs)
    else:
        frame = pd.read_excel(path, sheet_name=sheet or 0, header=None, skiprows=header_row - 1, engine="openpyxl")
    frame = frame.iloc[:, first:last].dropna(axis=1, how="all").dropna(axis=0, how="all")
    if frame.empty:
        raise ValueError("The selection contains no data. Choose a different header row or first column.")
    headers, used = [], {}
    for i, value in enumerate(frame.iloc[0]):
        text = str(value).strip() if pd.notna(value) else f"Column_{i + 1}"
        original = text
        index = 1
        while text in headers:
            index += 1
            text = f"{original}_{index}"
        used[original] = index
        headers.append(text)
    frame = frame.iloc[1:].reset_index(drop=True)
    frame.columns = headers
    if cache_dir:
        try:
            frame.to_parquet(cache_dir / f"{key}.parquet", index=False)
        except (ImportError, ValueError, TypeError):
            frame.to_pickle(cache_dir / f"{key}.pkl")
    return frame

def union_sources(frames: list[pd.DataFrame], mode: str = "stack", tag_column: str | None = None, tags: list[str] | None = None) -> pd.DataFrame:
    """Stack matching schemas and optionally retain source period tags."""
    if mode != "stack" or not frames:
        raise ValueError("Choose stack and at least one source.")
    expected = set(frames[0].columns)
    output = []
    for i, frame in enumerate(frames):
        if set(frame.columns) != expected:
            raise ValueError(f"Source {i + 1} has different columns. Missing: {sorted(expected - set(frame.columns))}; extra: {sorted(set(frame.columns) - expected)}.")
        frame = frame.copy()
        if tag_column:
            if tag_column in expected:
                raise ValueError(f"Tag column '{tag_column}' already exists. Choose another name.")
            frame[tag_column] = tags[i] if tags else str(i + 1)
        output.append(frame)
    return pd.concat(output, ignore_index=True)
